from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime


# ==========================================
# DATABASE PATH
# ==========================================

BASE_DIR = Path(__file__).resolve().parent.parent

EXCEL_FILE = (
    BASE_DIR /
    "database" /
    "FlowerBook_Data.xlsx"
)


# ==========================================
# GET WORKBOOK
# ==========================================

def get_workbook():

    return load_workbook(EXCEL_FILE)


# ==========================================
# GET FLOWERS
# ==========================================

def get_flowers():

    workbook = get_workbook()

    sheet = workbook["Flowers"]

    flowers = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        (
            flower_id,
            name,
            category,
            price,
            description,
            stock
        ) = row

        if flower_id is None:
            continue

        flowers.append({
            "id": flower_id,
            "name": name,
            "category": category,
            "price": price,
            "description": description,
            "stock": stock
        })

    workbook.close()

    return flowers


# ==========================================
# SAVE ORDER
# ==========================================

def save_order(order_data):

    workbook = get_workbook()

    customers_sheet = workbook["Customers"]
    orders_sheet = workbook["Orders"]

    customer_id = (
        f"CUST"
        f"{datetime.now().strftime('%Y%m%d%H%M%S')}"
    )

    order_id = (
        f"FB"
        f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
    )

    customers_sheet.append([
        customer_id,
        order_data["customer_name"],
        order_data["email"],
        order_data["phone"],
        order_data["address"]
    ])

    for item in order_data["items"]:

        orders_sheet.append([
            order_id,
            customer_id,
            order_data["customer_name"],
            order_data["email"],
            order_data["phone"],
            order_data["address"],
            item["name"],
            item["quantity"],
            item["price"] * item["quantity"],
            order_data["delivery_date"],
            order_data["delivery_time"],
            order_data.get("special_message", ""),
            "Confirmed"
        ])

    workbook.save(EXCEL_FILE)

    workbook.close()

    return {
        "order_id": order_id,
        "customer_id": customer_id,
        "status": "Confirmed"
    }


# ==========================================
# GET USER ORDERS
# ==========================================

def get_orders_by_email(email):

    workbook = get_workbook()

    sheet = workbook["Orders"]

    orders = {}

    normalized_email = email.strip().lower()


    # ==========================================
    # READ ORDERS
    # ==========================================

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        (
            order_id,
            customer_id,
            customer_name,
            order_email,
            phone,
            address,
            flower,
            quantity,
            amount,
            delivery_date,
            delivery_time,
            special_message,
            status
        ) = row


        if not order_id:
            continue


        if not order_email:
            continue


        # ==========================================
        # ONLY CURRENT USER
        # ==========================================

        if (
            str(order_email)
            .strip()
            .lower()
            != normalized_email
        ):

            continue


        # ==========================================
        # CREATE ORDER
        # ==========================================

        if order_id not in orders:

            orders[order_id] = {

                "order_id": order_id,

                "customer_id": customer_id,

                "customer_name": customer_name,

                "email": order_email,

                "phone": phone,

                "address": address,

                "delivery_date": delivery_date,

                "delivery_time": delivery_time,

                "special_message":
                    special_message,

                "status": status,

                "items": [],

                "total": 0

            }


        # ==========================================
        # ITEM DETAILS
        # ==========================================

        item_quantity = (
            int(quantity)
            if quantity is not None
            else 0
        )


        item_amount = (
            float(amount)
            if amount is not None
            else 0
        )


        unit_price = 0


        if item_quantity > 0:

            unit_price = (
                item_amount /
                item_quantity
            )


        orders[order_id]["items"].append({

            "name": flower,

            "quantity":
                item_quantity,

            "price":
                round(
                    unit_price,
                    2
                ),

            "amount":
                round(
                    item_amount,
                    2
                )

        })


        orders[order_id]["total"] += (
            item_amount
        )


    workbook.close()


    # ==========================================
    # FORMAT ORDERS
    # ==========================================

    order_list = []


    for order in orders.values():

        # Keep original item objects
        # for future reorder functionality.

        items_data = list(
            order["items"]
        )


        # ==========================================
        # FLOWER SUMMARY
        # ==========================================

        flower_names = []


        for item in items_data:

            flower_names.append(

                f'{item["name"]} × '
                f'{item["quantity"]}'

            )


        # Text shown on booking card
        order["items"] = ", ".join(
            flower_names
        )


        # Structured data for reorder
        order["itemsData"] = items_data


        # ==========================================
        # DELIVERY DATE
        # ==========================================

        if isinstance(
            order["delivery_date"],
            datetime
        ):

            order["delivery_date"] = (
                order["delivery_date"]
                .strftime("%Y-%m-%d")
            )


        # ==========================================
        # DELIVERY TIME
        # ==========================================

        if hasattr(
            order["delivery_time"],
            "strftime"
        ):

            order["delivery_time"] = (
                order["delivery_time"]
                .strftime("%H:%M")
            )


        # ==========================================
        # TOTAL
        # ==========================================

        order["total"] = round(
            order["total"],
            2
        )


        order_list.append(
            order
        )


    # ==========================================
    # NEWEST ORDERS FIRST
    # ==========================================

    order_list.reverse()


    return order_list