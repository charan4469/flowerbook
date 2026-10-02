from openpyxl import Workbook, load_workbook
from pathlib import Path


BASE_DIR = Path(__file__).resolve().parent.parent
DATABASE_DIR = BASE_DIR / "database"
EXCEL_FILE = DATABASE_DIR / "FlowerBook_Data.xlsx"


def create_database():
    DATABASE_DIR.mkdir(exist_ok=True)

    if EXCEL_FILE.exists():
        print("Excel database already exists.")
        return

    workbook = Workbook()

    # Customers sheet
    customers = workbook.active
    customers.title = "Customers"

    customers.append([
        "Customer ID",
        "Name",
        "Email",
        "Phone",
        "Address"
    ])

    # Orders sheet
    orders = workbook.create_sheet("Orders")

    orders.append([
        "Order ID",
        "Customer ID",
        "Customer Name",
        "Email",
        "Phone",
        "Address",
        "Flower",
        "Quantity",
        "Amount",
        "Delivery Date",
        "Delivery Time",
        "Special Message",
        "Status"
    ])

    # Flowers sheet
    flowers = workbook.create_sheet("Flowers")

    flowers.append([
        "Flower ID",
        "Name",
        "Category",
        "Price",
        "Description",
        "Stock"
    ])

    # Add initial flowers
    flowers.append([
        1,
        "Red Rose Bouquet",
        "Roses",
        599,
        "Beautiful fresh red roses arranged in an elegant bouquet.",
        50
    ])

    flowers.append([
        2,
        "Pink Lily Bouquet",
        "Lilies",
        799,
        "Fresh pink lilies perfect for special occasions.",
        40
    ])

    flowers.append([
        3,
        "Sunflower Delight",
        "Sunflowers",
        499,
        "Bright and cheerful sunflowers for someone special.",
        60
    ])

    workbook.save(EXCEL_FILE)

    print(f"Database created successfully: {EXCEL_FILE}")


if __name__ == "__main__":
    create_database()